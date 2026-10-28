#!/usr/bin/env python3
"""
Genspark PPT Exporter - CLI Tool
Converts Genspark slide projects and raw slides JSON into ultra-high-definition PPTX, PDF, and interactive HTML.
"""

import argparse
import asyncio
import json
import os
import sys
import shutil
from pathlib import Path
from PIL import Image

try:
    from playwright.async_api import async_playwright
except ImportError:
    print("Error: 'playwright' is not installed. Please run: pip install playwright && playwright install chromium")
    sys.exit(1)

try:
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:
    print("Error: 'python-pptx' is not installed. Please run: pip install python-pptx")
    sys.exit(1)


HTML_VIEWER_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{title} - Presentation</title>
  <style>
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{ background: #0f1117; color: #e2e8f0; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; display: flex; flex-direction: column; height: 100vh; overflow: hidden; }}
    header {{ background: #1a1f2c; padding: 10px 20px; display: flex; align-items: center; justify-content: space-between; border-bottom: 1px solid #2d3748; z-index: 10; }}
    .title {{ font-size: 15px; font-weight: 600; color: #f8fafc; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; max-width: 60%; }}
    .controls {{ display: flex; align-items: center; gap: 10px; }}
    button {{ background: #2d3748; color: #e2e8f0; border: 1px solid #4a5568; border-radius: 6px; padding: 6px 14px; font-size: 13px; font-weight: 500; cursor: pointer; transition: all 0.2s; }}
    button:hover {{ background: #4a5568; color: #fff; }}
    .counter {{ font-size: 13px; color: #a0aec0; min-width: 70px; text-align: center; font-variant-numeric: tabular-nums; }}
    main {{ flex: 1; display: flex; align-items: center; justify-content: center; position: relative; background: #000; overflow: hidden; }}
    iframe {{ width: 100%; height: 100%; border: none; }}
    .fullscreen-btn {{ background: #3b82f6; border-color: #2563eb; color: #fff; }}
    .fullscreen-btn:hover {{ background: #2563eb; }}
  </style>
</head>
<body>
  <header>
    <div class="title">{title}</div>
    <div class="controls">
      <button id="prevBtn" onclick="prevSlide()">◀ Previous</button>
      <span class="counter" id="counter">1 / {slide_count}</span>
      <button id="nextBtn" onclick="nextSlide()">Next ▶</button>
      <button class="fullscreen-btn" onclick="toggleFullScreen()">Full Screen (F)</button>
    </div>
  </header>
  <main id="slideContainer">
    <iframe id="slideFrame" src="{first_slide}"></iframe>
  </main>

  <script>
    const slides = {slides_json};
    let currentIndex = 0;
    const frame = document.getElementById('slideFrame');
    const counter = document.getElementById('counter');

    function updateSlide() {
      frame.src = slides[currentIndex];
      counter.textContent = `${currentIndex + 1} / ${slides.length}`;
    }

    function prevSlide() {
      if (currentIndex > 0) {
        currentIndex--;
        updateSlide();
      }
    }

    function nextSlide() {
      if (currentIndex < slides.length - 1) {
        currentIndex++;
        updateSlide();
      }
    }

    function toggleFullScreen() {
      if (!document.fullscreenElement) {
        document.getElementById('slideContainer').requestFullscreen().catch(err => {
          alert(`Full screen error: ${err.message}`);
        });
      } else {
        document.exitFullscreen();
      }
    }

    window.addEventListener('keydown', (e) => {
      if (e.key === 'ArrowRight' || e.key === ' ' || e.key === 'PageDown') {
        nextSlide();
      } else if (e.key === 'ArrowLeft' || e.key === 'PageUp') {
        prevSlide();
      } else if (e.key === 'f' || e.key === 'F') {
        toggleFullScreen();
      }
    });
  </script>
</body>
</html>
"""


async def render_slides(html_files, output_img_dir, scale=2.0):
    os.makedirs(output_img_dir, exist_ok=True)
    image_paths = []
    
    # Try system edge first, fall back to default chromium
    edge_paths = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
        "/usr/bin/microsoft-edge",
    ]
    exec_path = None
    for p in edge_paths:
        if os.path.exists(p):
            exec_path = p
            break

    async with async_playwright() as p:
        launch_kwargs = {"headless": True}
        if exec_path:
            launch_kwargs["executable_path"] = exec_path
            print(f"[*] Using browser: {exec_path}")
        else:
            print("[*] Using bundled Playwright Chromium")

        browser = await p.chromium.launch(**launch_kwargs)
        page = await browser.new_page(
            viewport={"width": 1920, "height": 1080},
            device_scale_factor=scale
        )

        for i, html_file in enumerate(html_files):
            idx = i + 1
            img_path = os.path.join(output_img_dir, f"slide_{idx:02d}.png")
            file_url = Path(html_file).resolve().as_uri()
            
            print(f"[{idx}/{len(html_files)}] Rendering {os.path.basename(html_file)}...")
            await page.goto(file_url, wait_until="networkidle")
            await page.wait_for_timeout(800)  # Wait for webfonts and animations to settle
            await page.screenshot(path=img_path)
            image_paths.append(img_path)

        await browser.close()
    return image_paths


def build_pptx(image_paths, output_pptx_path):
    print(f"[*] Assembling PowerPoint presentation: {output_pptx_path}...")
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for img_path in image_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)

    prs.save(output_pptx_path)
    print(f"[✓] PPTX successfully saved! Size: {os.path.getsize(output_pptx_path) / 1024 / 1024:.2f} MB")


def build_pdf(image_paths, output_pdf_path):
    print(f"[*] Generating PDF document: {output_pdf_path}...")
    pil_images = [Image.open(p).convert("RGB") for p in image_paths]
    if pil_images:
        pil_images[0].save(
            output_pdf_path,
            save_all=True,
            append_images=pil_images[1:],
            quality=95,
            resolution=150.0
        )
        print(f"[✓] PDF successfully saved! Size: {os.path.getsize(output_pdf_path) / 1024 / 1024:.2f} MB")


def export_from_json(json_path, output_dir, scale=2.0, formats=("pptx", "pdf", "html")):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    file_contents = data.get("data", {}).get("file_contents", [])
    if not file_contents and "file_contents" in data:
        file_contents = data["file_contents"]

    if not file_contents:
        print("Error: No slide contents found in the JSON file.")
        sys.exit(1)

    title = data.get("data", {}).get("meta_data", {}).get("title", "Genspark_Presentation")
    file_prefix = data.get("data", {}).get("meta_data", {}).get("file_prefix", "genspark_slides")
    # Clean file prefix
    file_prefix = "".join(c if c.isalnum() or c in ("-", "_") else "_" for c in file_prefix)

    html_dir = os.path.join(output_dir, "html")
    img_dir = os.path.join(output_dir, "images")
    os.makedirs(html_dir, exist_ok=True)
    os.makedirs(img_dir, exist_ok=True)

    html_files = []
    for i, item in enumerate(file_contents):
        idx = i + 1
        fn = item.get("filename", f"slide_{idx:02d}.html")
        content = item.get("content", "")
        fp = os.path.join(html_dir, fn)
        with open(fp, "w", encoding="utf-8") as f:
            f.write(content)
        html_files.append(fp)

    # Render images
    image_paths = asyncio.run(render_slides(html_files, img_dir, scale=scale))

    if "pptx" in formats:
        pptx_path = os.path.join(output_dir, f"{file_prefix}.pptx")
        build_pptx(image_paths, pptx_path)

    if "pdf" in formats:
        pdf_path = os.path.join(output_dir, f"{file_prefix}.pdf")
        build_pdf(image_paths, pdf_path)

    if "html" in formats:
        rel_html_paths = [f"html/{os.path.basename(f)}" for f in html_files]
        viewer_code = HTML_VIEWER_TEMPLATE.format(
            title=title,
            slide_count=len(html_files),
            first_slide=rel_html_paths[0] if rel_html_paths else "",
            slides_json=json.dumps(rel_html_paths)
        )
        viewer_path = os.path.join(output_dir, "index.html")
        with open(viewer_path, "w", encoding="utf-8") as f:
            f.write(viewer_code)
        print(f"[✓] Interactive HTML viewer saved to: {viewer_path}")

    print("\n[🎉] All done! Presentation exports completed in:", output_dir)


def main():
    parser = argparse.ArgumentParser(description="Genspark Slides to PPTX / PDF Exporter")
    parser.add_argument("input", help="Path to genspark-slides-data.json or folder containing slide HTML files")
    parser.add_argument("-o", "--output", default="./output", help="Output directory (default: ./output)")
    parser.add_argument("-s", "--scale", type=float, default=2.0, help="Device scale factor for rendering (default: 2.0 = 4K crispness)")
    parser.add_argument("-f", "--format", default="all", help="Output formats comma-separated: pptx,pdf,html or 'all' (default: all)")

    args = parser.parse_args()

    formats = ("pptx", "pdf", "html") if args.format == "all" else [x.strip().lower() for x in args.format.split(",")]
    
    input_path = os.path.abspath(args.input)
    if os.path.isfile(input_path):
        export_from_json(input_path, os.path.abspath(args.output), scale=args.scale, formats=formats)
    elif os.path.isdir(input_path):
        # Scan for HTML files
        html_files = sorted([os.path.join(input_path, f) for f in os.listdir(input_path) if f.endswith(".html")])
        if not html_files:
            print(f"Error: No .html files found in {input_path}")
            sys.exit(1)
        img_dir = os.path.join(args.output, "images")
        image_paths = asyncio.run(render_slides(html_files, img_dir, scale=args.scale))
        if "pptx" in formats:
            build_pptx(image_paths, os.path.join(args.output, "presentation.pptx"))
        if "pdf" in formats:
            build_pdf(image_paths, os.path.join(args.output, "presentation.pdf"))
    else:
        print(f"Error: Input path '{input_path}' not found.")
        sys.exit(1)


if __name__ == "__main__":
    main()
